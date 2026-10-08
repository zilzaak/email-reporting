import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    C_BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    C_BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    C_CARD_BG = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 232, 240)
    C_CARD_DARK = RGBColor(30, 41, 59)     # Slate 800
    
    C_PRIMARY = RGBColor(30, 64, 175)      # DIU Navy/Blue 800
    C_ACCENT_BLUE = RGBColor(37, 99, 235)  # Blue 600
    C_TEAL = RGBColor(13, 148, 136)        # Teal 600
    C_GREEN = RGBColor(16, 185, 129)       # Emerald 500
    C_AMBER = RGBColor(217, 119, 6)        # Amber 600
    C_RED = RGBColor(225, 29, 72)          # Rose 600
    C_PURPLE = RGBColor(124, 58, 237)      # Violet 600
    
    C_TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
    C_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    C_TEXT_LIGHT = RGBColor(255, 255, 255)
    C_TEXT_LIGHT_MUTED = RGBColor(203, 213, 225) # Slate 300
    
    blank_layout = prs.slide_layouts[6]
    
    def add_header(slide, title, category="GMAIL DELAY & SLA REPORTING SYSTEM"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_ACCENT_BLUE
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.73), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = C_CARD_BORDER
        line.line.fill.background()

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def draw_node(slide, x, y, w, h, icon, title, desc, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, title_color=C_TEXT_DARK):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        tb = slide.shapes.add_textbox(x + Inches(0.12), y + Inches(0.1), w - Inches(0.24), h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = title_color
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_MUTED if bg_color != C_CARD_DARK else C_TEXT_LIGHT_MUTED
        p2.space_before = Pt(3)
        return shape

    def draw_arrow_h(slide, x, y, length=Inches(0.4)):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y - Inches(0.1), length, Inches(0.2))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = C_ACCENT_BLUE
        arrow.line.fill.background()
        return arrow

    def draw_arrow_v(slide, x, y, length=Inches(0.4)):
        arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x - Inches(0.1), y, Inches(0.2), length)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = C_ACCENT_BLUE
        arrow.line.fill.background()
        return arrow

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_BG_DARK)
    
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.1), Inches(3.6), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    badge.line.fill.background()
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "DAFFODIL INTERNATIONAL UNIVERSITY"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(147, 197, 253)
    p_b.alignment = PP_ALIGN.CENTER
    
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.3), Inches(1.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "Gmail Delay & SLA Reporting System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_LIGHT
    
    p2 = tf_t.add_paragraph()
    p2.text = "Complete Mechanism, Architecture & Dual-Scheduler Workflow"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(10)

    card1 = create_card(s1, Inches(1.0), Inches(3.65), Inches(11.3), Inches(2.9), bg_color=C_CARD_DARK, border_color=RGBColor(51, 65, 85))
    tb_c = s1.shapes.add_textbox(Inches(1.3), Inches(3.85), Inches(10.7), Inches(2.5))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    
    p = tf_c.paragraphs[0]
    p.text = "PROJECT OVERVIEW & SPECIFICATIONS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_ACCENT_BLUE
    
    bullets = [
        "Core Engine: Two synchronized Spring Boot Schedulers (EmailMetadataFetchJob & EmailsReplierInfoFetchJob)",
        "Integration: Google Workspace REST API with Domain-Wide Delegation (DWD) Service Account impersonation",
        "Domain Monitored: @daffodilvarsity.edu.bd (Tracking HR Office, Registrar, Head Office, Accounts)",
        "Database Architecture: Microsoft SQL Server (Table: UM_HR_ESR_Email_Metadata, SP: SP_UM_HR_ESR_Email_Metadata_Save)",
        "Performance Optimization: Solved API redundancy via in-memory thread caching (70%+ API reduction)"
    ]
    for b in bullets:
        pb = tf_c.add_paragraph()
        pb.text = "•  " + b
        pb.font.size = Pt(13)
        pb.font.color.rgb = C_TEXT_LIGHT_MUTED
        pb.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: Business Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_BG_LIGHT)
    add_header(s2, "Business Problem & Chairman's Mandate", "Strategic Context")
    
    card_w = Inches(3.64)
    card_h = Inches(5.3)
    
    create_card(s2, Inches(0.8), Inches(1.65), card_w, card_h)
    tb1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "THE OPERATIONAL GAP"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_RED
    
    for pt in [
        "Communication Delays: Critical university offices issue instructions, circulars, and urgent administrative requests.",
        "Zero Visibility: University management previously had no automated way to determine who replied promptly vs. who delayed for days.",
        "Manual Audit Impossible: Thousands of employees across multiple campuses made manual mailbox checks completely impractical."
    ]:
        pp = tf1.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(12)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(12)

    create_card(s2, Inches(4.84), Inches(1.65), card_w, card_h)
    tb2 = s2.shapes.add_textbox(Inches(5.04), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "CHAIRMAN'S MANDATE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    
    for pt in [
        "Track Priority Mailboxes: Monitor all outbound communications from essential university departments:",
        "  - hr@daffodilvarsity.edu.bd\n  - registrar@daffodilvarsity.edu.bd\n  - accounts@, head office, IT desk",
        "SLA Standard: Response target of 8 hours during business operations.",
        "Delinquency Reporting: Daily reports showing delinquent response times with exact hours of delay."
    ]:
        pp = tf2.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(12)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(10)

    create_card(s2, Inches(8.88), Inches(1.65), card_w, card_h)
    tb3 = s2.shapes.add_textbox(Inches(9.08), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "AUTOMATED SOLUTION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_GREEN
    
    for pt in [
        "Zero User Friction: Cloud Service Account with Domain-Wide Delegation; zero employee passwords or plugins required.",
        "Two-Tier Engine: Scheduler #1 captures root threads & multi-recipients; Scheduler #2 dynamically captures delayed replies.",
        "Enterprise Persistence: Seamlessly integrates into core ERP database (ERPNewLatest) via stored procedures."
    ]:
        pp = tf3.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(12)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(12)

    # ==========================================
    # SLIDE 3: Technical Ecosystem
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_BG_LIGHT)
    add_header(s3, "Technical Ecosystem & Infrastructure Setup", "Architecture Foundation")
    
    tech_w = Inches(5.66)
    tech_h = Inches(2.55)
    
    create_card(s3, Inches(0.8), Inches(1.65), tech_w, tech_h)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), tech_w - Inches(0.4), tech_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "GOOGLE WORKSPACE & CLOUD DWD"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    for b in [
        "Domain-Wide Delegation (DWD): Service account authorized by IT Admin to impersonate approved university mailboxes.",
        "JSON Key: email-reporting-key.json managed via GmailFactory.",
        "Read-Only Scopes: https://www.googleapis.com/auth/gmail.readonly (Strict read-only access guarantees zero risk of data mutation)."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(4)

    create_card(s3, Inches(6.86), Inches(1.65), tech_w, tech_h)
    tb = s3.shapes.add_textbox(Inches(7.06), Inches(1.8), tech_w - Inches(0.4), tech_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "JAVA & SPRING BOOT BACKEND"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_ACCENT_BLUE
    for b in [
        "Framework: Spring Boot 3 running on Java 17/21 runtime.",
        "Core Services: EmailMetadataFetchService & EmailMetadataPersistToDBService.",
        "Scheduling: Spring @Scheduled cron engines running asynchronously without blocking core application threads.",
        "Timezone: Strictly configured for Asia/Dhaka (UTC+6)."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(4)

    create_card(s3, Inches(0.8), Inches(4.45), tech_w, tech_h)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(4.6), tech_w - Inches(0.4), tech_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MICROSOFT SQL SERVER & PERSISTENCE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    for b in [
        "Database: ERPNewLatest on SQL Server 2019 Dialect.",
        "Target Table: UM_HR_ESR_Email_Metadata (atomic row per recipient).",
        "Procedure: SP_UM_HR_ESR_Email_Metadata_Save with dual operation modes (I = Insert, U = Update).",
        "Spring Data JPA & JdbcTemplate with SimpleJdbcCall."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(4)

    create_card(s3, Inches(6.86), Inches(4.45), tech_w, tech_h)
    tb = s3.shapes.add_textbox(Inches(7.06), Inches(4.6), tech_w - Inches(0.4), tech_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MANAGED MAILBOX FLEET"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_GREEN
    for b in [
        "Approved Mailboxes (Configured in application.properties):",
        "  - hr@daffodilvarsity.edu.bd (Central HR Office)",
        "  - registrar@daffodilvarsity.edu.bd (Office of Registrar)",
        "  - hroffice4@, hroffice9@, hroffice10@ (HR Officers)",
        "  - ithelpdesk@daffodilvarsity.edu.bd (IT Helpdesk)",
        "Configurable Lookback: 1 day lookback (configurable)."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: FLOWCHART 1 - Date & Mailbox Unique Thread Retrieval Engine
    # =========================================================================
    sf1 = prs.slides.add_slide(blank_layout)
    set_slide_background(sf1, C_BG_LIGHT)
    add_header(sf1, "FLOWCHART 1: Date & Mailbox Unique Thread Retrieval", "Heart #1 Flowchart (Part A)")
    
    # 5 Sequential Step Cards across top/middle
    col_w = Inches(2.1)
    col_h = Inches(3.6)
    spacing = Inches(0.3)
    start_x = Inches(0.8)
    y_pos = Inches(1.8)
    
    fc1_nodes = [
        ("⏰ [Cron Trigger]", "EmailMetadataFetchJob", "Triggers daily/periodic.\nSets targetDate = yesterday.\nTimezone: Asia/Dhaka.", C_PRIMARY),
        ("🗓️ [Epoch Math]", "Compute Bounds", "start = target.atStartOfDay\nend = plusDays(1).atStartOfDay\nBoundary: Dec 31 23:59:59 to Jan 2 00:00:00.", C_TEAL),
        ("🔍 [Gmail Query]", "Build Search Query", "after:{start-1} before:{end}\nfrom:{mailbox}\n(e.g., hr@daffodilvarsity.edu.bd)", C_ACCENT_BLUE),
        ("🌐 [Google API]", "threads.list()", "Iterates approved mailboxes.\nmaxResults = 100L.\nPaginates with pageToken.\nReturns UNIQUE thread IDs.", C_PURPLE),
        ("📥 [Fetch Metadata]", "threads.get()", "Fetches thread details:\nformat = 'metadata'\nHeaders: From, To, Cc, Subject, Date.\n(Full tree, zero date clip)", C_GREEN)
    ]
    
    for i, (icon, title, desc, col) in enumerate(fc1_nodes):
        x = start_x + i * (col_w + spacing)
        draw_node(sf1, x, y_pos, col_w, col_h, icon, title, desc, border_color=col, title_color=col)
        if i < 4:
            draw_arrow_h(sf1, x + col_w + Inches(0.05), y_pos + Inches(1.8), length=Inches(0.2))

    # Bottom Decision / Clarification Box
    bot_card = create_card(sf1, Inches(0.8), Inches(5.7), Inches(11.73), Inches(1.4), bg_color=C_CARD_DARK, border_color=RGBColor(51, 65, 85))
    tb_b = sf1.shapes.add_textbox(Inches(1.1), Inches(5.8), Inches(11.1), Inches(1.2))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "💡 CRITICAL LOGICAL GUARANTEE: WHAT IF NO ONE REPLIED ON JAN 1?"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    
    p2 = tf_b.add_paragraph()
    p2.text = (
        "• Google's threads.list(q) checks if ANY message in the thread matches 'from:mailbox' and the date boundary.\n"
        "• The initial email (Message 0) was sent by the mailbox on Jan 1. Therefore, Google ALWAYS returns the thread ID.\n"
        "• Unreplied threads are 100% reliably captured and passed to Flowchart 2 for multi-recipient row insertion."
    )
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_TEXT_LIGHT_MUTED
    p2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 5: FLOWCHART 2 - Root Dissection & Multi-Recipient DB Insertion
    # =========================================================================
    sf2 = prs.slides.add_slide(blank_layout)
    set_slide_background(sf2, C_BG_LIGHT)
    add_header(sf2, "FLOWCHART 2: Multi-Recipient Breakdown & Atomic Row Insert", "Heart #1 Flowchart (Part B)")
    
    # 3-Tier Flowchart Layout
    # Tier 1: Root extraction
    t1_w = Inches(3.6)
    t1_h = Inches(1.7)
    draw_node(sf2, Inches(0.8), Inches(1.7), t1_w, t1_h, "✉️ [Root Message]", "Extract messages.get(0)", 
              "Sender = hr@daffodilvarsity.edu.bd\nTo-List = [emp1@, emp2@, emp3@]\nDelivered = 2026-09-01 09:00:00", 
              border_color=C_PRIMARY, title_color=C_PRIMARY)
    
    draw_arrow_h(sf2, Inches(4.5), Inches(2.5), length=Inches(0.3))
    
    # Tier 2: Reply detection
    draw_node(sf2, Inches(4.9), Inches(1.7), Inches(3.8), t1_h, "🔍 [Scan messages 1..n]", "Match 'From' against To-List", 
              "If reply sender in receiver list:\n  delay = (replyEpoch - rootEpoch) / 3.6e6\nKeep EARLIEST reply if multiple replies.", 
              border_color=C_AMBER, title_color=C_AMBER)
    
    draw_arrow_h(sf2, Inches(8.8), Inches(2.5), length=Inches(0.3))
    
    # Tier 3: Bulk String Serialization
    draw_node(sf2, Inches(9.2), Inches(1.7), Inches(3.3), t1_h, "📦 [Array Serialization]", "Vectorized Parameters", 
              "bulkEmailTo: 'emp1@,emp2@,emp3@'\nbulkReplyDate: '10:15,NULL,14:00'\nbulkDelay: '1.25,NULL,5.00'", 
              border_color=C_PURPLE, title_color=C_PURPLE)

    # Down arrow to Stored Procedure
    draw_arrow_v(sf2, Inches(6.6), Inches(3.55), length=Inches(0.35))

    # Stored Procedure Node
    sp_box = create_card(sf2, Inches(2.5), Inches(4.0), Inches(8.33), Inches(0.75), bg_color=RGBColor(238, 242, 255), border_color=C_PRIMARY)
    tb_sp = sf2.shapes.add_textbox(Inches(2.6), Inches(4.05), Inches(8.13), Inches(0.65))
    tf_sp = tb_sp.text_frame
    p = tf_sp.paragraphs[0]
    p.text = "⚡ STORED PROCEDURE: SP_UM_HR_ESR_Email_Metadata_Save (Operation = 'I')"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf_sp.add_paragraph()
    p2.text = "Splits comma-separated lists and generates atomic discrete records for all N recipients linked by root_thread_id"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_DARK
    p2.alignment = PP_ALIGN.CENTER

    draw_arrow_v(sf2, Inches(6.6), Inches(4.85), length=Inches(0.35))

    # Database Output Rows (3 discrete cards)
    row_w = Inches(3.64)
    row_h = Inches(1.7)
    
    draw_node(sf2, Inches(0.8), Inches(5.3), row_w, row_h, "✅ [Row 1: Replied]", "emp1@daffodilvarsity.edu.bd", 
              "delivered_date: 09:00:00\nreply_date: 10:15:00\nresponse_delay_hour: 1.25\nstatus: 1 (true / completed)", 
              border_color=C_GREEN, title_color=C_GREEN)
    
    draw_node(sf2, Inches(4.84), Inches(5.3), row_w, row_h, "⏳ [Row 2: PENDING]", "emp2@daffodilvarsity.edu.bd", 
              "delivered_date: 09:00:00\nreply_date: NULL\nresponse_delay_hour: NULL\nstatus: 0 (false / pending reply)", 
              border_color=C_RED, title_color=C_RED)
    
    draw_node(sf2, Inches(8.88), Inches(5.3), row_w, row_h, "✅ [Row 3: Replied]", "emp3@daffodilvarsity.edu.bd", 
              "delivered_date: 09:00:00\nreply_date: 14:00:00\nresponse_delay_hour: 5.00\nstatus: 1 (true / completed)", 
              border_color=C_GREEN, title_color=C_GREEN)

    # =========================================================================
    # SLIDE 6: FLOWCHART 3 - Scheduler 2 Unreplied Tracking & Cache Sync
    # =========================================================================
    sf3 = prs.slides.add_slide(blank_layout)
    set_slide_background(sf3, C_BG_LIGHT)
    add_header(sf3, "FLOWCHART 3: Scheduler #2 Unreplied Sync & Cache Optimization", "Heart #2 Flowchart")
    
    # Left: Database Query -> Cache Check -> Google API
    c3_w = Inches(3.6)
    c3_h = Inches(2.2)
    
    draw_node(sf3, Inches(0.8), Inches(1.7), c3_w, c3_h, "🗄️ [Step 1: DB Query]", "Fetch Pending Rows", 
              "SELECT id, rootThreadId,\n       receiverEmail, senderEmail\nFROM EmailMetadata\nWHERE status = false\nORDER BY deliveredDate ASC", 
              border_color=C_PRIMARY, title_color=C_PRIMARY)
    
    draw_arrow_h(sf3, Inches(4.5), Inches(2.8), length=Inches(0.3))
    
    draw_node(sf3, Inches(4.9), Inches(1.7), Inches(3.8), c3_h, "⚡ [Step 2: replyCache Check]", "Prevent Duplicate Calls", 
              "Map<String, EmailMetadataDTO> replyCache\n\nIs rootThreadId in cache?\n• YES ➜ Reuse DTO (0 API calls!)\n• NO  ➜ Fetch Google API once & put in cache", 
              border_color=C_GREEN, title_color=C_GREEN)
    
    draw_arrow_h(sf3, Inches(8.8), Inches(2.8), length=Inches(0.3))
    
    draw_node(sf3, Inches(9.2), Inches(1.7), Inches(3.3), c3_h, "🌐 [Step 3: Google Fetch]", "threads.get() by rootThreadId", 
              "Impersonates senderEmail.\nReturns FULL message history.\nCaptures all replies sent since Day 1.", 
              border_color=C_ACCENT_BLUE, title_color=C_ACCENT_BLUE)

    draw_arrow_v(sf3, Inches(6.6), Inches(4.0), length=Inches(0.35))

    # Decision Box
    dec_box = create_card(sf3, Inches(2.5), Inches(4.45), Inches(8.33), Inches(0.95), bg_color=RGBColor(254, 243, 199), border_color=C_AMBER)
    tb_dec = sf3.shapes.add_textbox(Inches(2.6), Inches(4.5), Inches(8.13), Inches(0.85))
    tf_dec = tb_dec.text_frame
    p = tf_dec.paragraphs[0]
    p.text = "❓ DECISION: DID THE PENDING RECEIVER EMAIL REPLY IN THIS CONVERSATION?"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_AMBER
    p.alignment = PP_ALIGN.CENTER
    p2 = tf_dec.add_paragraph()
    p2.text = "• NO  ➜ Leave row in DB as status=0 (pending). Will re-check on next cycle.\n• YES ➜ Extract replyDate, replyMessageId, calculate delayHour = (replyTime - rootTime) / 3.6e6"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_DARK
    p2.space_before = Pt(2)
    p2.alignment = PP_ALIGN.CENTER

    draw_arrow_v(sf3, Inches(6.6), Inches(5.5), length=Inches(0.35))

    # Batch Update Box
    end_box = create_card(sf3, Inches(1.5), Inches(5.95), Inches(10.33), Inches(1.15), bg_color=C_CARD_DARK, border_color=RGBColor(51, 65, 85))
    tb_end = sf3.shapes.add_textbox(Inches(1.7), Inches(6.05), Inches(9.93), Inches(0.95))
    tf_end = tb_end.text_frame
    p = tf_end.paragraphs[0]
    p.text = "🚀 STEP 4: BATCH STORED PROCEDURE UPDATE (Operation = 'U')"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p2 = tf_end.add_paragraph()
    p2.text = (
        "Accumulates bulkIds, bulkReplyDates, bulkReplyDelayHours ➜ calls saveToDB(..., 'U').\n"
        "SQL Server updates target rows in a single batch query, transitions status = 1 (true), and completes SLA tracking!"
    )
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = C_TEXT_LIGHT_MUTED
    p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 7: High-Level Architecture
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_BG_LIGHT)
    add_header(s7, "High-Level Architecture: The Two-Scheduler Pipeline", "System Design")
    
    flow_w = Inches(3.64)
    flow_h = Inches(4.9)
    
    create_card(s7, Inches(0.8), Inches(1.7), flow_w, flow_h, border_color=C_PRIMARY)
    badge1 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), flow_w, Inches(0.5))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = C_PRIMARY
    badge1.line.fill.background()
    p = badge1.text_frame.paragraphs[0]
    p.text = "HEART #1: SNAPSHOT ENGINE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.35), flow_w - Inches(0.4), flow_h - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "EmailMetadataFetchJob"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    for b in [
        "Frequency: Scheduled nightly / recurring.",
        "Target: Yesterday's completed date.",
        "Inspects: Approved mailbox senders.",
        "Calls: Google threads.list & threads.get.",
        "Identifies: Root message (index 0).",
        "Inserts: 1 row per recipient into DB.",
        "Initial State: Replied employees get delay hours; pending employees get NULL."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(6)

    create_card(s7, Inches(4.84), Inches(1.7), flow_w, flow_h, border_color=C_TEAL)
    badge2 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.84), Inches(1.7), flow_w, Inches(0.5))
    badge2.fill.solid()
    badge2.fill.fore_color.rgb = C_TEAL
    badge2.line.fill.background()
    p = badge2.text_frame.paragraphs[0]
    p.text = "CENTRAL DATA STORE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    
    tb = s7.shapes.add_textbox(Inches(5.04), Inches(2.35), flow_w - Inches(0.4), flow_h - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "UM_HR_ESR_Email_Metadata"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    for b in [
        "Relational Table: Linked by rootThreadId.",
        "Partitioning: One row per (rootThreadId, receiverEmail).",
        "status = 1 (true): Recipient has replied.",
        "status = 0 (false): Recipient pending reply (awaiting Scheduler 2).",
        "Stored Procedure: SP_UM_HR_ESR_Email_Metadata_Save ensures ACID compliance for bulk operations."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(8)

    create_card(s7, Inches(8.88), Inches(1.7), flow_w, flow_h, border_color=C_ACCENT_BLUE)
    badge3 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.88), Inches(1.7), flow_w, Inches(0.5))
    badge3.fill.solid()
    badge3.fill.fore_color.rgb = C_ACCENT_BLUE
    badge3.line.fill.background()
    p = badge3.text_frame.paragraphs[0]
    p.text = "HEART #2: SLA SYNC ENGINE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    
    tb = s7.shapes.add_textbox(Inches(9.08), Inches(2.35), flow_w - Inches(0.4), flow_h - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "EmailsReplierInfoFetchJob"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    for b in [
        "Frequency: Periodic (cron = '0 0/2 * * * ?').",
        "Queries DB: Fetches unreplied rows (status=false).",
        "Optimized Fetch: Caches by rootThreadId to avoid duplicate API calls.",
        "Scans Threads: Detects new replies from recipients.",
        "Bulk Update: Calls Stored Procedure with operation 'U'.",
        "Outcome: Automatically closes delay tracking when replies arrive."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + b
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(6)

    # ==========================================
    # SLIDE 8: Gmail Thread Semantics Clarification
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, C_BG_LIGHT)
    add_header(s8, "Gmail Thread Architecture: Semantics & Proof", "API Deep Dive")
    
    col_w = Inches(5.66)
    col_h = Inches(5.3)
    
    create_card(s8, Inches(0.8), Inches(1.65), col_w, col_h)
    tb1 = s8.shapes.add_textbox(Inches(1.0), Inches(1.85), col_w - Inches(0.4), col_h - Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "CORE QUESTION: HOW THREADS BEHAVE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    for item in [
        "Developer's Critical Question:",
        "  'Is a thread unique over a specific conversation, or does threads.list also return reply-type threads?'",
        "The Technical Truth:",
        "  1. A Thread ID is 100% UNIQUE per conversation tree.",
        "  2. threads.list() returns THREAD IDs, NOT individual messages.",
        "  3. Even if 10 people reply and forward inside a conversation, it remains ONE SINGLE THREAD ID.",
        "  4. threads.list() returns that thread ID exactly ONCE."
    ]:
        pp = tf1.add_paragraph()
        pp.text = "• " + item if not item.startswith("  ") else item
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(8)

    create_card(s8, Inches(6.86), Inches(1.65), col_w, col_h)
    tb2 = s8.shapes.add_textbox(Inches(7.06), Inches(1.85), col_w - Inches(0.4), col_h - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "DOES threads.get() FILTER BY DATE?"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_GREEN
    for item in [
        "Definitive Answer: NO! threads.get() NEVER filters by date.",
        "What threads.get('me', threadId) Actually Returns:",
        "  It returns EVERY message ever sent in that conversation, from day 1 up to the current second.",
        "Why This Is The Key To Scheduler 2's Success:",
        "  - When Scheduler 1 ran on Day 1, only Message [0] existed.",
        "  - Recipient C replies on Day 5.",
        "  - When Scheduler 2 calls threads.get() on Day 5, it receives Message [0] AND Reply [1].",
        "  - Scheduler 2 detects Reply [1] immediately without needing complex date queries!"
    ]:
        pp = tf2.add_paragraph()
        pp.text = "• " + item if not item.startswith("  ") else item
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(8)

    # ==========================================
    # SLIDE 9: Redundancy Analysis & Optimization Solution
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, C_BG_LIGHT)
    add_header(s9, "Redundancy Analysis & Thread Caching Optimization", "Performance Optimization")
    
    create_card(s9, Inches(0.8), Inches(1.65), Inches(5.66), Inches(5.3))
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(5.26), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "THE REDUNDANCY HAZARD IDENTIFIED"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_RED
    for rp in [
        "Scenario: HR sends 1 notice to 5 employees (A, B, C, D, E). None replied yet.",
        "Database State: 5 distinct rows in UM_HR_ESR_Email_Metadata, all with rootThreadId = '18f3a2b'.",
        "Naive Loop in Scheduler 2:",
        "  Row 1 (A) -> calls threads.get('18f3a2b') [API Call 1]",
        "  Row 2 (B) -> calls threads.get('18f3a2b') [API Call 2]",
        "  Row 3 (C) -> calls threads.get('18f3a2b') [API Call 3]",
        "  Row 4 (D) -> calls threads.get('18f3a2b') [API Call 4]",
        "  Row 5 (E) -> calls threads.get('18f3a2b') [API Call 5]",
        "Problem: 5 calls to Google API for the EXACT SAME conversation thread! Wastes quota and slows performance."
    ]:
        pp = tf.add_paragraph()
        pp.text = "• " + rp if not rp.startswith("  Row") else rp
        pp.font.size = Pt(11)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(6)

    create_card(s9, Inches(6.86), Inches(1.65), Inches(5.66), Inches(5.3), bg_color=C_CARD_DARK, border_color=RGBColor(51,65,85))
    tb_r = s9.shapes.add_textbox(Inches(7.06), Inches(1.85), Inches(5.26), Inches(4.9))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "OUR SOLUTION: IN-MEMORY THREAD CACHE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    
    sol_code = (
        "// Implemented in EmailMetadataPersistToDBService\n"
        "Map<String, EmailMetadataDTO> replyCache = new HashMap<>();\n\n"
        "for (Map<String, Object> mp : notRepliedThreadList) {\n"
        "    String threadId = (String) mp.get(\"rootThreadId\");\n\n"
        "    if (replyCache.containsKey(threadId)) {\n"
        "        // REUSE CACHED THREAD METADATA (0 API CALLS)\n"
        "        dto = replyCache.get(threadId);\n"
        "    } else {\n"
        "        // FETCH ONCE PER UNIQUE THREAD\n"
        "        dto = fetchService.processThread(gmail, threadId, ...);\n"
        "        replyCache.put(threadId, dto);\n"
        "    }\n"
        "    // Evaluate receiver on cached dto...\n"
        "}"
    )
    pp = tf_r.add_paragraph()
    pp.text = sol_code
    pp.font.size = Pt(9.5)
    pp.font.name = "Consolas"
    pp.font.color.rgb = RGBColor(226, 232, 240)
    pp.space_before = Pt(10)
    
    badge_opt = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.06), Inches(5.9), Inches(5.26), Inches(0.75))
    badge_opt.fill.solid()
    badge_opt.fill.fore_color.rgb = RGBColor(6, 78, 59)
    badge_opt.line.color.rgb = C_GREEN
    p = badge_opt.text_frame.paragraphs[0]
    p.text = "PROVEN IMPACT: 70%+ API Call Reduction! 1,000 pending rows across 250 threads now takes 250 API calls instead of 1,000."
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(167, 243, 208)
    p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 10: Edge Cases Handled in Production
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, C_BG_LIGHT)
    add_header(s10, "Edge Cases & Production Resilience", "Reliability Engineering")
    
    grid_w = Inches(3.64)
    grid_h = Inches(2.45)
    
    cases = [
        ("Multiple Replies from Same User", "Earliest Reply Rule: When an employee sends several follow-ups, the shortest delay timestamp is strictly preserved.", C_PRIMARY),
        ("Sender Replies to Own Thread", "Recipient Validation: hr@ sending a follow-up is ignored because hr@ is not in the extracted receiverEmails list.", C_ACCENT_BLUE),
        ("Forwarded Conversations", "Root Isolation: Even if someone forwards the email externally, index 0 remains the original root sender.", C_TEAL),
        ("Large Recipient Lists", "Bulk Comma Parsing: Handles emails sent to 50+ recipients using vectorized array packing and stored procedure splitting.", C_AMBER),
        ("Mailbox Access Exceptions", "Fault Tolerance: Isolated try-catch around individual mailboxes and threads ensures one timeout never crashes the scheduler.", C_RED),
        ("Timezone Boundary Drift", "Asia/Dhaka Epoch Alignment: Hardened epoch math guarantees zero date skew between UTC server time and local office time.", C_GREEN)
    ]
    
    for i, (title, desc, color) in enumerate(cases):
        row = i // 3
        col = i % 3
        x = Inches(0.8 + col * 4.04)
        y = Inches(1.65 + row * 2.7)
        card = create_card(s10, x, y, grid_w, grid_h)
        tb = s10.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), grid_w - Inches(0.4), grid_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        pp = tf.add_paragraph()
        pp.text = desc
        pp.font.size = Pt(10.5)
        pp.font.color.rgb = C_TEXT_DARK
        pp.space_before = Pt(6)

    # ==========================================
    # SLIDE 11: Summary & Q&A
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, C_BG_DARK)
    
    t_box = s11.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.33), Inches(1.5))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "Summary & Project Manager Q&A"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_LIGHT
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf_t.add_paragraph()
    p2.text = "Ready to Address All Questions on Architecture, Accuracy & Performance"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)

    rc = create_card(s11, Inches(1.5), Inches(2.9), Inches(10.33), Inches(3.8), bg_color=C_CARD_DARK, border_color=RGBColor(51, 65, 85))
    tb_rc = s11.shapes.add_textbox(Inches(1.8), Inches(3.1), Inches(9.73), Inches(3.4))
    tf_rc = tb_rc.text_frame
    tf_rc.word_wrap = True
    p = tf_rc.paragraphs[0]
    p.text = "EXECUTIVE TAKEAWAYS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_ACCENT_BLUE
    
    recap = [
        "1. Core Requirement Fulfilled: Automatic daily tracking of employee email response delays from critical offices.",
        "2. Thread Mechanics Proven: Gmail threads are single conversation entities; threads.get() retrieves the entire thread without date clipping.",
        "3. Zero-Reply Handling Validated: Unreplied emails are correctly captured and written to the database with NULL reply dates.",
        "4. Redundancy Defeated: In-memory thread caching (replyCache) delivers a 70%+ reduction in Google API calls.",
        "5. Production-Ready: Backed by MS SQL Server stored procedures and robust exception handling."
    ]
    for r in recap:
        pp = tf_rc.add_paragraph()
        pp.text = r
        pp.font.size = Pt(12)
        pp.font.color.rgb = C_TEXT_LIGHT_MUTED
        pp.space_before = Pt(8)

    output_path = os.path.abspath("Gmail_SLA_Delay_Reporting_System.pptx")
    try:
        prs.save(output_path)
        print(f"Presentation saved successfully to: {output_path}")
    except PermissionError:
        output_path_v2 = os.path.abspath("Gmail_SLA_Delay_Reporting_Flowcharts.pptx")
        prs.save(output_path_v2)
        print(f"File locked by PowerPoint. Saved updated presentation with flowcharts to: {output_path_v2}")

if __name__ == "__main__":
    create_deck()
